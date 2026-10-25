from datetime import timedelta
from io import StringIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook

from .models import Exchange, SharedActivity, SheetSource, Shift, SyncState
from .sheets import sync_workbook


class RepeatedResidentImportTests(TestCase):
    def setUp(self):
        self.day = (timezone.localtime() + timedelta(days=30)).replace(
            hour=0, minute=0, second=0, microsecond=0, tzinfo=None
        )
        self.source = 'repeated-resident-regression'

    def workbook(self, first='Pessoa Teste', second=None):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Espelho Semanal'
        sheet.append([None, self.day])
        sheet.append([None, 'Segunda'])
        sheet.append(['INSTITUTO', 'Ambulatório'])
        sheet.append([None, 'R1', '8-12', None, None, first, second])
        return book

    def test_same_resident_in_both_columns_is_one_shift(self):
        book = self.workbook(second='Pessoa Teste')
        sync_workbook(book, self.source)
        sync_workbook(book, self.source)
        self.assertEqual(Shift.objects.filter(source=self.source, active=True).count(), 1)
        self.assertFalse(SharedActivity.objects.filter(source=self.source).exists())

    def test_case_and_spacing_variations_are_not_a_pair(self):
        sync_workbook(self.workbook(second='PESSOA   TESTE'), self.source)
        self.assertEqual(Shift.objects.filter(source=self.source, active=True).count(), 1)
        self.assertFalse(SharedActivity.objects.filter(source=self.source).exists())

    def old_duplicate(self):
        # Two separate rows simulate the two occurrence IDs produced by the old
        # F/G loop, while ensuring separate source rows remain supported.
        book = self.workbook()
        book.active.append([None, 'R1', '8-12', None, None, 'Pessoa Teste'])
        sync_workbook(book, self.source)
        shifts = list(Shift.objects.filter(source=self.source).order_by('pk'))
        self.assertEqual(len(shifts), 2)
        book.active.delete_rows(5)
        book.active.cell(row=4, column=7, value='Pessoa Teste')
        return book, shifts

    def test_reimport_deactivates_old_duplicate_and_preserves_first_id(self):
        book, (first, duplicate) = self.old_duplicate()
        sync_workbook(book, self.source)
        first.refresh_from_db()
        duplicate.refresh_from_db()
        self.assertTrue(first.active)
        self.assertFalse(duplicate.active)
        self.assertEqual(Shift.objects.filter(source=self.source, active=True).count(), 1)

    def test_reimport_protects_an_accepted_duplicate_for_reconciliation(self):
        book, (first, duplicate) = self.old_duplicate()
        original_owner = duplicate.resident
        other_owner = get_user_model().objects.create_user('other-duplicate-owner')
        other = Shift.objects.create(
            external_id='other-duplicate-shift', resident=original_owner,
            starts_at=duplicate.starts_at + timedelta(days=1),
            ends_at=duplicate.ends_at + timedelta(days=1), sector='Outro setor',
        )
        duplicate.resident = other_owner
        duplicate.save(update_fields=['resident'])
        Exchange.objects.create(
            offered=duplicate, requested=other, proposer=original_owner,
            recipient=other_owner, status='accepted', resolved_at=timezone.now(),
        )
        summary = sync_workbook(book, self.source)
        duplicate.refresh_from_db()
        self.assertTrue(duplicate.active)
        self.assertEqual(duplicate.resident_id, other_owner.pk)
        self.assertIn('1 exigem conciliação', summary)

    def test_bad_layout_reports_source_line_and_preserves_existing_data(self):
        source_id = 'c' * 30
        book = self.workbook()
        sync_workbook(book, source_id)
        before = list(Shift.objects.filter(source=source_id).values())
        last_success = SyncState.objects.get(source=source_id).last_success
        book.active.cell(row=3, column=2).value = None
        stderr = StringIO()
        with patch('schedule.management.commands.sync_sheets.fetch_workbook', return_value=book):
            with self.assertRaises(CommandError):
                call_command('sync_sheets', source_id, only=True, stderr=stderr)
        self.assertEqual(list(Shift.objects.filter(source=source_id).values()), before)
        self.assertEqual(SyncState.objects.get(source=source_id).last_success, last_success)
        error = SheetSource.objects.get(spreadsheet_id=source_id).error
        self.assertIn('Linha 4', error)
        self.assertIn('setor não identificado', error)
        self.assertIn('dados anteriores preservados', stderr.getvalue())

    def test_unknown_error_never_exposes_exception_details(self):
        source_id = 'd' * 30
        stderr = StringIO()
        error = ValueError('https://example.invalid/export?key=private-test-value')
        with patch('schedule.management.commands.sync_sheets.fetch_workbook', side_effect=error):
            with self.assertRaises(CommandError):
                call_command('sync_sheets', source_id, only=True, stderr=stderr)
        message = SheetSource.objects.get(spreadsheet_id=source_id).error
        self.assertIn('ValueError', message)
        self.assertNotIn('private-test-value', message + stderr.getvalue())
        self.assertNotIn('example.invalid', message + stderr.getvalue())
