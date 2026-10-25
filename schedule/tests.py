from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.utils import timezone
from .models import Shift, Exchange
from .services import propose_exchange, respond_exchange

class WorkflowTests(TestCase):
    def setUp(self):
        self.a = get_user_model().objects.create_user('ana')
        self.b = get_user_model().objects.create_user('bia')
        self.c = get_user_model().objects.create_user('carla')
        self.start = (timezone.now()+timedelta(days=7)).replace(microsecond=0)
        self.s1 = Shift.objects.create(external_id='1',resident=self.a,starts_at=self.start,ends_at=self.start+timedelta(hours=6),sector='UTI')
        self.s2 = Shift.objects.create(external_id='2',resident=self.b,starts_at=self.start+timedelta(days=1),ends_at=self.start+timedelta(days=1,hours=6),sector='PS')
    def propose(self):
        return propose_exchange(self.a,self.s1.pk,self.s2.pk)
    def test_requires_login(self):
        self.assertEqual(self.client.get('/').status_code,302)
    def test_proposal_does_not_change_owners(self):
        self.assertEqual(self.propose().status,'pending')
        self.s1.refresh_from_db()
        self.assertEqual(self.s1.resident,self.a)
    def test_both_parties_required(self):
        exchange=self.propose()
        for user in [self.a,self.c]:
            with self.assertRaises(ValidationError): respond_exchange(user,exchange.pk,'accept')
        respond_exchange(self.b,exchange.pk,'accept')
        self.s1.refresh_from_db(); self.s2.refresh_from_db()
        self.assertEqual((self.s1.resident,self.s2.resident),(self.b,self.a))
        with self.assertRaises(ValidationError): respond_exchange(self.b,exchange.pk,'accept')
    def test_cannot_offer_someone_elses_shift(self):
        with self.assertRaises(ValidationError): propose_exchange(self.c,self.s1.pk,self.s2.pk)
    def test_reject_preserves_owners(self):
        exchange=self.propose()
        respond_exchange(self.b,exchange.pk,'reject')
        self.s1.refresh_from_db()
        self.assertEqual(self.s1.resident,self.a)
    def test_cancel(self):
        exchange=self.propose()
        self.assertEqual(respond_exchange(self.a,exchange.pk,'cancel').status,'cancelled')
    def test_conflict_detected_at_acceptance(self):
        exchange=self.propose()
        Shift.objects.create(external_id='3',resident=self.a,starts_at=self.s2.starts_at,ends_at=self.s2.ends_at,sector='PS')
        with self.assertRaises(ValidationError): respond_exchange(self.b,exchange.pk,'accept')
        self.s1.refresh_from_db()
        self.assertEqual(self.s1.resident,self.a)
    def test_third_party_cannot_access_exchange(self):
        exchange=self.propose()
        self.client.force_login(self.c)
        self.assertEqual(self.client.post(f'/trocas/{exchange.pk}/accept/').status_code,404)
    def test_dashboard_and_csrf(self):
        self.client.force_login(self.a)
        self.assertContains(self.client.get('/'),'Minha escala')
        from django.test import Client
        client=Client(enforce_csrf_checks=True)
        client.force_login(self.b)
        self.assertEqual(client.post('/trocas/criar/',{}).status_code,403)
    def test_import_repeatable_preserves_exchange(self):
        with TemporaryDirectory() as folder:
            path=Path(folder)/'schedule.csv'
            path.write_text(f'id,residente,inicio,fim,setor\n1,ana,{self.s1.starts_at.isoformat()},{self.s1.ends_at.isoformat()},UTI\n2,bia,{self.s2.starts_at.isoformat()},{self.s2.ends_at.isoformat()},PS\n')
            call_command('import_schedule',str(path))
            call_command('import_schedule',str(path))
            self.assertEqual(Shift.objects.count(),2)
            exchange=self.propose()
            respond_exchange(self.b,exchange.pk,'accept')
            call_command('import_schedule',str(path))
            self.s1.refresh_from_db()
            self.assertEqual(self.s1.resident,self.b)
    def test_invalid_import_rolls_back(self):
        from django.core.management.base import CommandError
        with TemporaryDirectory() as folder:
            path=Path(folder)/'bad.csv'
            path.write_text(f'id,residente,inicio,fim,setor\nnew,ana,{self.s1.starts_at.isoformat()},{self.s1.ends_at.isoformat()},UTI\nbad,unknown,invalid,invalid,PS\n')
            with self.assertRaises(CommandError): call_command('import_schedule',str(path))
            self.assertEqual(Shift.objects.count(),2)
    def test_interest_is_owner_only(self):
        self.client.force_login(self.b)
        self.assertEqual(self.client.post(f'/plantoes/{self.s1.pk}/interesse/').status_code,404)
        self.client.force_login(self.a)
        self.assertEqual(self.client.post(f'/plantoes/{self.s1.pk}/interesse/').status_code,302)
        self.s1.refresh_from_db()
        self.assertTrue(self.s1.wants_exchange)
    def test_sheet_sync_idempotent_and_removed_shift(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import RotationActivity
        book=Workbook(); sheet=book.active;sheet.title='Espelho Semanal'
        day=self.start.replace(tzinfo=None)
        sheet.append([None,day]);sheet.append([None,'Segunda'])
        sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','R1-A',None,'Residente Teste'])
        rotation=book.create_sheet('R1')
        rotation.append(['R1',None,'Período',day,day+timedelta(days=7)])
        rotation.append(['Segunda','Terça','Quarta','Quinta','Sexta'])
        rotation.append([None]);rotation.append([None])
        rotation.append(['R1 - A','Residente Teste'])
        rotation.append(['AMB']*5);rotation.append(['CC']*5)
        sync_workbook(book,'test-source');sync_workbook(book,'test-source')
        self.assertEqual(Shift.objects.filter(source='test-source').count(),1)
        self.assertEqual(RotationActivity.objects.count(),0)
        imported=Shift.objects.get(source='test-source')
        self.assertFalse(imported.resident.has_usable_password())
        imported.wants_exchange=True;imported.save()
        sync_workbook(book,'test-source');imported.refresh_from_db()
        self.assertTrue(imported.wants_exchange)
        sheet.cell(row=4,column=6,value='Outro Residente')
        sync_workbook(book,'test-source');imported.refresh_from_db()
        self.assertFalse(imported.active)
    def test_past_and_inactive_cannot_be_exchanged(self):
        self.s1.active=False;self.s1.save()
        with self.assertRaises(ValidationError): self.propose()
    def test_acceptance_clears_interest(self):
        self.s1.wants_exchange=True;self.s1.save()
        respond_exchange(self.b,self.propose().pk,'accept')
        self.s1.refresh_from_db()
        self.assertFalse(self.s1.wants_exchange)
    def test_shared_names_are_alternatives_not_two_shifts(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import SharedActivity
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)])
        sheet.append([None,'Segunda'])
        sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','C1A-H','C1B-A','Pessoa Um','Pessoa Dois'])
        sheet.append(['EXTERNO','Wet lab'])
        sheet.append([None,'R1','10-12','C1A-H','C1B-A','Pessoa Um','Pessoa Dois'])
        sync_workbook(book,'paired-source')
        self.assertEqual(Shift.objects.filter(source='paired-source',active=True).count(),0)
        self.assertEqual(SharedActivity.objects.filter(source='paired-source').count(),2)
        for activity in SharedActivity.objects.filter(source='paired-source'):
            self.assertEqual(activity.residents.count(),2)
        sync_workbook(book,'paired-source')
        self.assertEqual(SharedActivity.objects.filter(source='paired-source').count(),2)
    def allocation_pair(self):
        from .models import SharedActivity
        start=self.start+timedelta(days=10)
        mine=SharedActivity.objects.create(external_id='shared-one',source='test',starts_at=start,ends_at=start+timedelta(hours=6),sector='Ambulatório')
        theirs=SharedActivity.objects.create(external_id='shared-two',source='test',starts_at=start,ends_at=start+timedelta(hours=3),sector='Wet lab')
        mine.residents.set([self.a,self.b]);theirs.residents.set([self.a,self.b])
        return mine,theirs
    def test_allocation_updates_calendar_only_after_both_confirm(self):
        from .services import propose_allocation,respond_allocation
        mine,theirs=self.allocation_pair()
        allocation=propose_allocation(self.a,mine.pk,theirs.pk)
        self.assertEqual(Shift.objects.count(),2)
        with self.assertRaises(ValidationError): respond_allocation(self.a,allocation.pk,'accept')
        respond_allocation(self.b,allocation.pk,'accept')
        self.assertEqual(Shift.objects.count(),4)
        mine.refresh_from_db();theirs.refresh_from_db()
        self.assertEqual((mine.allocated_to_id,theirs.allocated_to_id),(self.a.pk,self.b.pk))
        self.client.force_login(self.a)
        self.assertContains(self.client.get('/calendario/',{'month':mine.starts_at.strftime('%Y-%m')}),'Ambulatório')
        ics=self.client.get('/calendario.ics').content.decode()
        self.assertIn('SUMMARY:Ambulatório',ics)
        self.assertNotIn('SUMMARY:Wet lab',ics)
        with self.assertRaises(ValidationError): respond_allocation(self.b,allocation.pk,'accept')
    def test_allocation_cannot_be_proposed_by_third_party(self):
        from .services import propose_allocation
        mine,theirs=self.allocation_pair()
        with self.assertRaises(ValidationError): propose_allocation(self.c,mine.pk,theirs.pk)
    def test_rejected_allocation_creates_no_calendar_entries(self):
        from .services import propose_allocation,respond_allocation
        mine,theirs=self.allocation_pair()
        allocation=propose_allocation(self.a,mine.pk,theirs.pk)
        respond_allocation(self.b,allocation.pk,'reject')
        self.assertEqual(Shift.objects.count(),2)
    def test_allocation_conflict_rolls_back_both_assignments(self):
        from .services import propose_allocation,respond_allocation
        mine,theirs=self.allocation_pair()
        allocation=propose_allocation(self.a,mine.pk,theirs.pk)
        Shift.objects.create(external_id='conflict',resident=self.b,starts_at=theirs.starts_at,ends_at=theirs.ends_at,sector='Outro')
        with self.assertRaises(ValidationError): respond_allocation(self.b,allocation.pk,'accept')
        mine.refresh_from_db();theirs.refresh_from_db()
        self.assertIsNone(mine.allocated_to_id);self.assertIsNone(theirs.allocated_to_id)
        self.assertEqual(Shift.objects.count(),3)
    def test_allocation_api_protects_third_party(self):
        from .services import propose_allocation
        mine,theirs=self.allocation_pair()
        allocation=propose_allocation(self.a,mine.pk,theirs.pk)
        self.client.force_login(self.c)
        self.assertEqual(self.client.post(f'/divisoes/{allocation.pk}/accept/').status_code,404)
    def test_sheets_preserve_mutually_confirmed_allocation(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import SharedActivity
        from .services import propose_allocation,respond_allocation
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,(self.start+timedelta(days=20)).replace(tzinfo=None)])
        sheet.append([None,'Segunda']);sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','X','Y','Pessoa Um','Pessoa Dois'])
        sheet.append(['EXTERNO','Wet lab'])
        sheet.append([None,'R1','10-12','X','Y','Pessoa Um','Pessoa Dois'])
        sync_workbook(book,'confirm-source')
        mine,theirs=list(SharedActivity.objects.filter(source='confirm-source'))
        a,b=list(mine.residents.all())
        proposal=propose_allocation(a,mine.pk,theirs.pk)
        respond_allocation(b,proposal.pk,'accept')
        sync_workbook(book,'confirm-source')
        mine.refresh_from_db();theirs.refresh_from_db()
        self.assertEqual((mine.allocated_to_id,theirs.allocated_to_id),(a.pk,b.pk))
        self.assertEqual(Shift.objects.filter(external_id__startswith='allocation:',active=True).count(),2)
    def test_linked_administrator_uses_resident_identity(self):
        from .models import ResidentAccess
        admin=get_user_model().objects.create_superuser('coordenador',password='test-only-password')
        self.client.force_login(admin)
        response=self.client.post('/conta/vincular/',{'resident':self.a.pk})
        self.assertEqual(response.status_code,302)
        self.assertEqual(ResidentAccess.objects.get(account=admin).resident,self.a)
        response=self.client.get('/')
        self.assertTrue(response.context['has_offered'])
        self.assertEqual(response.context['resident_user'],self.a)
        self.client.post('/trocas/criar/',{'offered':self.s1.pk,'requested':self.s2.pk})
        self.assertEqual(Exchange.objects.get().proposer,self.a)
        ics=self.client.get('/calendario.ics').content.decode()
        self.assertIn(f'UID:shift-{self.s1.pk}@escala',ics)
        self.assertNotIn(f'UID:shift-{self.s2.pk}@escala',ics)
    def test_resident_cannot_link_to_another_identity(self):
        from .models import ResidentAccess
        self.client.force_login(self.a)
        self.assertEqual(self.client.post('/conta/vincular/',{'resident':self.b.pk}).status_code,403)
        self.assertFalse(ResidentAccess.objects.exists())
    def test_no_future_shifts_explains_why_exchange_is_disabled(self):
        Shift.objects.update(starts_at=timezone.now()-timedelta(days=2),ends_at=timezone.now()-timedelta(days=1))
        self.client.force_login(self.a)
        response=self.client.get('/')
        self.assertContains(response,'Não há plantões futuros importados')
        self.assertContains(response,'disabled')
        self.assertNotContains(response,'<h2>Referências de rotação</h2>')
    def test_headers_in_time_column_and_missing_times(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import IncompleteActivity
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)])
        sheet.append([None,'Segunda'])
        sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','A',None,'Pessoa Um'])
        sheet.append(['UNEX','R2','U.E. Centro Norte'])
        sheet.append([None,None,'9-18','B',None,'Pessoa Dois'])
        sheet.append([None,None,None,'C',None,'Pessoa Tres'])
        sync_workbook(book,'header-source')
        shift=Shift.objects.get(source='header-source',source_cell='C6:G6')
        self.assertEqual(shift.sector,'UNEX · U.E. Centro Norte')
        self.assertEqual(timezone.localtime(shift.starts_at).hour,9)
        self.assertIn('09:00',str(shift))
        self.assertEqual(IncompleteActivity.objects.get(source='header-source').source_cell,'C7:G7')
        self.assertEqual(Shift.objects.filter(source='header-source').count(),2)
    def test_linked_recipient_can_confirm_exchange(self):
        from .models import ResidentAccess
        exchange=self.propose()
        account=get_user_model().objects.create_superuser('admin-recipient',password='test-only-password')
        ResidentAccess.objects.create(account=account,resident=self.b)
        self.client.force_login(account)
        self.assertContains(self.client.get('/'),f'/trocas/{exchange.pk}/accept/')
        self.client.post(f'/trocas/{exchange.pk}/accept/')
        exchange.refresh_from_db()
        self.assertEqual(exchange.status,'accepted')
    def test_drive_discovers_month_subfolders_and_all_pages(self):
        from unittest.mock import Mock
        from .drive import discover_folder,FOLDER_MIME,SHEET_MIME
        from .models import DriveFolder,SheetSource
        folder=DriveFolder.objects.create(folder_id='rootfolder0001',name='Escalas')
        SheetSource.objects.create(spreadsheet_id='spreadsheet002',enabled=False)
        payloads=[
            {'files':[{'id':'monthfolder001','name':'Outubro','mimeType':FOLDER_MIME}],'nextPageToken':'page-two'},
            {'files':[{'id':'spreadsheet001','name':'Semana 1','mimeType':SHEET_MIME}]},
            {'files':[{'id':'spreadsheet002','name':'Semana 2','mimeType':SHEET_MIME}]},
        ]
        session=Mock()
        responses=[]
        for payload in payloads:
            response=Mock();response.json.return_value=payload;responses.append(response)
        session.get.side_effect=responses
        self.assertEqual(discover_folder(session,folder),2)
        self.assertEqual(SheetSource.objects.count(),2)
        self.assertFalse(SheetSource.objects.get(spreadsheet_id='spreadsheet002').enabled)
        self.assertEqual(session.get.call_args_list[1].kwargs['params']['pageToken'],'page-two')
        self.assertIn('monthfolder001',session.get.call_args_list[2].kwargs['params']['q'])
    def test_failed_drive_discovery_does_not_save_partial_listing(self):
        from unittest.mock import Mock
        from .drive import discover_folder,FOLDER_MIME,SHEET_MIME
        from .models import DriveFolder,SheetSource
        folder=DriveFolder.objects.create(folder_id='rootfolder0001')
        response=Mock();response.json.return_value={'files':[{'id':'spreadsheet001','name':'Semana','mimeType':SHEET_MIME},{'id':'monthfolder001','name':'Mês','mimeType':FOLDER_MIME}]}
        session=Mock();session.get.side_effect=[response,RuntimeError('Unavailable')]
        with self.assertRaises(RuntimeError): discover_folder(session,folder)
        self.assertFalse(SheetSource.objects.exists())
        folder.refresh_from_db();self.assertIsNone(folder.last_success)
    def test_multiple_weekly_sources_are_unified_without_erasing_previous_week(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        for offset,source in [(0,'week-source-one'),(7,'week-source-two')]:
            book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
            sheet.append([None,(self.start+timedelta(days=offset)).replace(tzinfo=None)])
            sheet.append([None,'Segunda']);sheet.append(['INSTITUTO','Ambulatório'])
            sheet.append([None,'R1','8-18','A',None,'Pessoa Um'])
            sync_workbook(book,source)
            if offset==7: sync_workbook(book,source)
        self.assertEqual(Shift.objects.filter(source__startswith='week-source',active=True).count(),2)
        self.assertEqual(Shift.objects.filter(source__startswith='week-source').values('resident_id').distinct().count(),1)
    def test_dashboard_opens_nearest_upcoming_week(self):
        self.client.force_login(self.a)
        response=self.client.get('/')
        day=timezone.localtime(self.s1.starts_at).date()
        self.assertEqual(response.context['week_start'],day-timedelta(days=day.weekday()))
        self.assertContains(response,'Solicitar troca')
    def test_week_selection_is_preserved_explicitly(self):
        self.client.force_login(self.a)
        response=self.client.get('/',{'week':'2026-09-28'})
        self.assertEqual(response.context['week_start'].isoformat(),'2026-09-28')
    def test_source_registration_is_administrator_only(self):
        self.client.force_login(self.a)
        self.assertEqual(self.client.get('/?tab=sources').status_code,403)
        self.assertEqual(self.client.post('/fontes/adicionar/',{'url':'https://docs.google.com/spreadsheets/d/'+'a'*30+'/edit'}).status_code,403)
    def test_admin_can_add_a_week_without_duplicates(self):
        from .models import SheetSource
        account=get_user_model().objects.create_superuser('admin-sources',password='test-only-password')
        self.client.force_login(account)
        url='https://docs.google.com/spreadsheets/d/'+'a'*30+'/edit?usp=sharing'
        for _ in range(2):
            self.assertEqual(self.client.post('/fontes/adicionar/',{'url':url,'name':'Próxima semana'}).status_code,302)
        self.assertEqual(SheetSource.objects.count(),1)
        self.assertContains(self.client.get('/?tab=sources'),'Próxima semana')
        self.assertContains(self.client.get('/?tab=sources'),'Aguardando importação')
    def test_source_registration_rejects_other_hosts(self):
        from .models import SheetSource
        account=get_user_model().objects.create_superuser('admin-invalid-source',password='test-only-password')
        self.client.force_login(account)
        for url in ['https://example.com/spreadsheets/d/'+'a'*30+'/edit','http://docs.google.com/spreadsheets/d/'+'a'*30+'/edit','https://docs.google.com.evil.example/spreadsheets/d/'+'a'*30+'/edit']:
            self.client.post('/fontes/adicionar/',{'url':url})
        self.assertEqual(SheetSource.objects.count(),0)
    def test_application_assets_are_registered(self):
        from django.contrib.staticfiles import finders
        self.assertTrue(finders.find('schedule/app.css'))
        self.assertTrue(finders.find('schedule/app.js'))
    def test_sync_labels_an_added_week_from_its_actual_dates(self):
        from openpyxl import Workbook
        from unittest.mock import patch
        from .models import SheetSource
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        from datetime import datetime
        sheet.append([None,datetime(2026,10,26)])
        sheet.append([None,'Segunda']);sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','A',None,'Pessoa Um'])
        with patch('schedule.management.commands.sync_sheets.fetch_workbook',return_value=book):
            call_command('sync_sheets','a'*30)
        self.assertEqual(SheetSource.objects.get().name,'Semana de 26/10 a 01/11/2026')
    def test_interest_publishes_one_topic_and_redirects_to_forum(self):
        from .models import ExchangeTopic
        from .services import set_exchange_interest
        self.client.force_login(self.a)
        response=self.client.post(f'/plantoes/{self.s1.pk}/interesse/')
        topic=ExchangeTopic.objects.get()
        self.assertRedirects(response,f'/mural/{topic.pk}/')
        self.assertEqual(topic.poster,self.a)
        self.assertEqual(set_exchange_interest(self.a,self.s1.pk,True).pk,topic.pk)
        self.assertEqual(ExchangeTopic.objects.count(),1)
        self.client.force_login(self.b)
        self.assertContains(self.client.get('/mural/'),'UTI')
        self.assertContains(self.client.get(f'/mural/{topic.pk}/'),'Propor uma troca')
    def test_forum_requires_authentication(self):
        self.assertEqual(self.client.get('/mural/').status_code,302)
    def test_forum_comments_are_visible_and_html_is_escaped(self):
        from .services import set_exchange_interest
        from .models import ForumComment
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        self.client.force_login(self.b)
        text='<script>alert(1)</script> Consigo trocar na quarta.'
        response=self.client.post(f'/mural/{topic.pk}/comentar/',{'text':text,'author':self.c.pk})
        self.assertEqual(response.status_code,302)
        self.assertEqual(ForumComment.objects.get().author,self.b)
        response=self.client.get(f'/mural/{topic.pk}/')
        self.assertContains(response,'&lt;script&gt;alert(1)&lt;/script&gt;')
        self.assertNotContains(response,'<script>alert(1)</script>')
        self.client.force_login(self.c)
        self.assertContains(self.client.get(f'/mural/{topic.pk}/'),'Consigo trocar na quarta.')
    def test_blank_and_oversized_comments_are_rejected(self):
        from .services import set_exchange_interest,add_forum_comment
        from .models import ForumComment
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        for text in ['   ','x'*601]:
            with self.assertRaises(ValidationError): add_forum_comment(self.b,topic.pk,text)
        self.assertEqual(ForumComment.objects.count(),0)
    def test_withdraw_is_owner_only_and_preserves_comments(self):
        from .services import set_exchange_interest,add_forum_comment,withdraw_forum_topic
        from .models import ForumComment
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        add_forum_comment(self.b,topic.pk,'Tenho interesse.')
        self.client.force_login(self.b)
        self.assertEqual(self.client.post(f'/mural/{topic.pk}/retirar/').status_code,404)
        withdraw_forum_topic(self.a,topic.pk);withdraw_forum_topic(self.a,topic.pk)
        self.s1.refresh_from_db();topic.refresh_from_db()
        self.assertFalse(self.s1.wants_exchange);self.assertIsNotNone(topic.closed_at)
        self.assertEqual(ForumComment.objects.count(),1)
        with self.assertRaises(ValidationError): add_forum_comment(self.b,topic.pk,'Novo comentário')
        self.client.force_login(self.a)
        self.assertContains(self.client.get('/mural/?mine=1'),'Encerrado')
        self.assertNotContains(self.client.get('/mural/'),f'href="/mural/{topic.pk}/"')
    def test_accepted_exchange_closes_both_forum_topics(self):
        from .services import set_exchange_interest
        from .models import ExchangeTopic
        first=set_exchange_interest(self.a,self.s1.pk,True)
        second=set_exchange_interest(self.b,self.s2.pk,True)
        respond_exchange(self.b,self.propose().pk,'accept')
        self.assertEqual(ExchangeTopic.objects.filter(closed_at__isnull=False).count(),2)
        first.refresh_from_db();second.refresh_from_db()
        self.assertFalse(first.is_open);self.assertFalse(second.is_open)
        self.client.force_login(self.c)
        response=self.client.get('/mural/')
        self.assertNotContains(response,f'href="/mural/{first.pk}/"')
        self.assertNotContains(response,f'href="/mural/{second.pk}/"')
    def test_stale_or_expired_forum_topics_cannot_receive_comments(self):
        from .services import set_exchange_interest,add_forum_comment
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        self.s1.active=False;self.s1.save()
        with self.assertRaises(ValidationError): add_forum_comment(self.b,topic.pk,'Posso trocar')
        self.client.force_login(self.b)
        self.assertNotContains(self.client.get('/mural/'),f'href="/mural/{topic.pk}/"')
    def test_forum_uses_linked_resident_for_comments(self):
        from .services import set_exchange_interest
        from .models import ResidentAccess,ForumComment
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        admin=get_user_model().objects.create_superuser('forum-admin',password='test-only-password')
        ResidentAccess.objects.create(account=admin,resident=self.b)
        self.client.force_login(admin)
        self.client.post(f'/mural/{topic.pk}/comentar/',{'text':'Quero combinar.'})
        self.assertEqual(ForumComment.objects.get().author,self.b)
    def test_existing_interests_can_be_migrated_to_topics_without_duplicates(self):
        from importlib import import_module
        from django.apps import apps
        from django.db import connection
        from unittest.mock import Mock
        from .models import ExchangeTopic
        self.s1.wants_exchange=True;self.s1.save()
        migration=import_module('schedule.migrations.0010_seed_exchange_topics')
        editor=Mock();editor.connection.alias=connection.alias
        migration.seed_topics(apps,editor);migration.seed_topics(apps,editor)
        self.assertEqual(ExchangeTopic.objects.count(),1)
        self.assertEqual(ExchangeTopic.objects.get().poster,self.a)
    def test_new_owner_gets_new_forum_thread_after_exchange(self):
        from .services import set_exchange_interest,add_forum_comment
        from .models import ForumComment
        original=set_exchange_interest(self.a,self.s1.pk,True)
        add_forum_comment(self.c,original.pk,'Conversa anterior.')
        respond_exchange(self.b,self.propose().pk,'accept')
        new=set_exchange_interest(self.b,self.s1.pk,True)
        self.assertNotEqual(new.pk,original.pk)
        self.assertEqual(new.poster,self.b)
        self.assertEqual(ForumComment.objects.filter(topic=new).count(),0)
        self.assertEqual(ForumComment.objects.filter(topic=original).count(),1)
    def test_failed_acceptance_keeps_forum_publication_open(self):
        from .services import set_exchange_interest
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        exchange=self.propose()
        Shift.objects.create(external_id='forum-conflict',resident=self.b,starts_at=self.s1.starts_at,ends_at=self.s1.ends_at,sector='Conflito')
        with self.assertRaises(ValidationError): respond_exchange(self.b,exchange.pk,'accept')
        topic.refresh_from_db();self.s1.refresh_from_db();exchange.refresh_from_db()
        self.assertTrue(topic.is_open)
        self.assertTrue(self.s1.wants_exchange)
        self.assertEqual(exchange.status,'pending')
    def test_reimport_preserves_an_existing_forum_thread(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .services import set_exchange_interest
        from .models import ExchangeTopic
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)])
        sheet.append([None,'Segunda']);sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','A',None,'Pessoa de teste'])
        sync_workbook(book,'forum-source')
        shift=Shift.objects.get(source='forum-source')
        topic=set_exchange_interest(shift.resident,shift.pk,True)
        sync_workbook(book,'forum-source')
        topic.refresh_from_db()
        self.assertTrue(topic.is_open)
        self.assertEqual(ExchangeTopic.objects.count(),1)
    def test_live_status_requires_login_and_contains_no_schedule_content(self):
        self.assertEqual(self.client.get('/dados/status/').status_code,302)
        self.client.force_login(self.a)
        response=self.client.get('/dados/status/')
        self.assertEqual(response.status_code,200)
        self.assertEqual(set(response.json()),{'revision'})
        self.assertEqual(len(response.json()['revision']),64)
        self.assertNotIn('ana',response.content.decode())
        self.assertEqual(response['Cache-Control'],'private, no-store')
    def test_live_revision_changes_after_publication_comment_and_confirmation(self):
        from .live import data_revision
        from .services import set_exchange_interest,add_forum_comment
        original=data_revision()
        self.assertEqual(original,data_revision())
        topic=set_exchange_interest(self.a,self.s1.pk,True)
        published=data_revision();self.assertNotEqual(original,published)
        add_forum_comment(self.b,topic.pk,'Posso trocar.')
        commented=data_revision();self.assertNotEqual(published,commented)
        respond_exchange(self.b,self.propose().pk,'accept')
        self.assertNotEqual(commented,data_revision())
    def test_public_drive_uses_key_file_without_service_account(self):
        from unittest.mock import patch
        from .drive import authenticated_session
        with TemporaryDirectory() as folder:
            path=Path(folder)/'api-key.txt';path.write_text('test-key-not-a-real-credential\n')
            with patch.dict('os.environ',{'GOOGLE_DRIVE_API_KEY_FILE':str(path),'GOOGLE_APPLICATION_CREDENTIALS':'/not-needed.json'}):
                session=authenticated_session()
                try:
                    self.assertTrue(session.public_drive)
                    self.assertEqual(session.params['key'],'test-key-not-a-real-credential')
                    self.assertTrue(session.verify)
                finally: session.close()
    def test_public_folder_sync_uses_public_sheet_export(self):
        from unittest.mock import patch,Mock
        from openpyxl import Workbook
        from .models import DriveFolder,SheetSource
        folder=DriveFolder.objects.create(folder_id='rootfolder0001')
        SheetSource.objects.create(spreadsheet_id='b'*30,folder=folder)
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)])
        sheet.append([None,'Segunda']);sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1','8-18','A',None,'Pessoa Um'])
        session=Mock();session.public_drive=True
        with patch('schedule.drive.authenticated_session',return_value=session),patch('schedule.drive.discover_folder',return_value=1),patch('schedule.management.commands.sync_sheets.fetch_workbook',return_value=book) as public_export,patch('schedule.drive.fetch_private_workbook') as private_export:
            call_command('sync_sheets',registered=True)
            public_export.assert_called_once_with('b'*30)
            private_export.assert_not_called()
        self.assertEqual(Shift.objects.filter(source='b'*30,active=True).count(),1)
    def test_midnight_end_is_the_following_day(self):
        from .sheets import parse_shift_hours
        day=timezone.localtime(self.start).date()
        start,end=parse_shift_hours(day,'12-24')
        self.assertEqual(timezone.localtime(start).hour,12)
        self.assertEqual(timezone.localtime(end).hour,0)
        self.assertEqual(timezone.localtime(end).date(),day+timedelta(days=1))
        self.assertIsNone(parse_shift_hours(day,'12-24:30'))
    def test_rows_without_names_never_create_missing_time_records(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import IncompleteActivity
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)]);sheet.append([None,'Segunda'])
        sheet.append(['INSTITUTO','Ambulatório'])
        sheet.append([None,'R1',None,None,None,None,None])
        sheet.append([None,None,'99-99',None,None,'   ',None])
        sheet.append([None,None,None,None,None,'Pessoa Um',None])
        sync_workbook(book,'blank-name-source')
        self.assertEqual(IncompleteActivity.objects.filter(source='blank-name-source').count(),1)
        self.assertEqual(IncompleteActivity.objects.get(source='blank-name-source').source_cell,'C6:G6')
    def test_location_marker_does_not_erase_the_sector_above_names(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import IncompleteActivity
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)]);sheet.append([None,'Segunda'])
        sheet.append(['OUTROS',None,'PÓS PLANTÃO R2'])
        sheet.append(['EXTERNO',None,None,None,None,'Pessoa Um'])
        summary=sync_workbook(book,'rest-source')
        activity=IncompleteActivity.objects.get(source='rest-source')
        self.assertTrue(activity.is_rest)
        self.assertIn('PÓS PLANTÃO R2',activity.sector)
        self.assertIn('0 sem horário válido; 1 registros de pós-plantão',summary)
    def test_targeted_retry_does_not_import_other_registered_sources(self):
        from unittest.mock import patch
        from openpyxl import Workbook
        from .models import SheetSource
        SheetSource.objects.create(spreadsheet_id='a'*30)
        SheetSource.objects.create(spreadsheet_id='b'*30)
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)]);sheet.append([None,'Segunda'])
        sheet.append(['INSTITUTO','Ambulatório']);sheet.append([None,'R1','12-24',None,None,'Pessoa Um'])
        with patch('schedule.management.commands.sync_sheets.fetch_workbook',return_value=book) as fetch:
            call_command('sync_sheets','a'*30,only=True)
            fetch.assert_called_once_with('a'*30)
    def test_sequential_activities_cannot_be_divided(self):
        from .services import propose_allocation
        mine,theirs=self.allocation_pair()
        theirs.starts_at=mine.ends_at+timedelta(hours=1)
        theirs.ends_at=theirs.starts_at+timedelta(hours=4);theirs.save()
        with self.assertRaisesMessage(ValidationError,'Atividades em sequência'):
            propose_allocation(self.a,mine.pk,theirs.pk)
        self.client.force_login(self.a)
        response=self.client.get('/')
        self.assertEqual(response.context['allocation_options'],[])
    def test_acceptance_revalidates_overlapping_times(self):
        from .services import propose_allocation,respond_allocation
        mine,theirs=self.allocation_pair()
        proposal=propose_allocation(self.a,mine.pk,theirs.pk)
        theirs.starts_at=mine.ends_at;theirs.ends_at=mine.ends_at+timedelta(hours=3);theirs.save()
        with self.assertRaisesMessage(ValidationError,'Atividades em sequência'):
            respond_allocation(self.b,proposal.pk,'accept')
        mine.refresh_from_db();self.assertIsNone(mine.allocated_to_id)
    def test_one_resident_in_two_rows_has_two_sequential_shifts(self):
        from openpyxl import Workbook
        from .sheets import sync_workbook
        from .models import SharedActivity
        book=Workbook();sheet=book.active;sheet.title='Espelho Semanal'
        sheet.append([None,self.start.replace(tzinfo=None)]);sheet.append([None,'Segunda'])
        sheet.append(['HOSPITAL','Sala Dr. Pedro']);sheet.append([None,None,'8-12',None,None,'Residente Um'])
        sheet.append(['INSTITUTO','Patologia']);sheet.append([None,None,'13-18',None,None,'Residente Um'])
        sync_workbook(book,'sequential-source')
        self.assertEqual(Shift.objects.filter(source='sequential-source',active=True).count(),2)
        self.assertFalse(SharedActivity.objects.filter(source='sequential-source').exists())
    def test_sector_color_is_stable_and_grouping_preserves_spreadsheet_order(self):
        from .presentation import sector_hue,sector_rows
        self.assertEqual(sector_hue('INSTITUTO · Refração'),sector_hue('instituto · refracao'))
        self.assertNotEqual(sector_hue('UTI'),sector_hue('Patologia'))
        self.s1.source_cell='C112:G112';self.s1.save()
        self.s2.source_cell='J11:N11';self.s2.save()
        groups=sector_rows([{'date':self.start.date(),'shifts':[self.s1,self.s2],'shared':[],'incomplete':[]}])
        self.assertEqual([g['name'] for g in groups],['PS','UTI'])
        self.assertEqual(sum(c['count'] for g in groups for c in g['cells']),2)
    def test_month_calendar_groups_multiple_shifts_in_the_same_sector(self):
        from .presentation import month_sector_groups
        self.s2.sector=self.s1.sector
        groups=month_sector_groups([self.s1,self.s2])
        self.assertEqual(len(groups),1);self.assertEqual(len(groups[0]['shifts']),2)
    def test_requests_inbox_only_exposes_requests_addressed_to_resident(self):
        incoming=self.propose()
        outgoing=Exchange.objects.create(offered=self.s2,requested=self.s1,proposer=self.b,recipient=self.a)
        self.client.force_login(self.b)
        received=self.client.get('/?tab=requests&box=received')
        self.assertContains(received,'Aceitar troca')
        self.assertEqual(list(received.context['exchanges']),[incoming])
        sent=self.client.get('/?tab=requests&box=sent')
        self.assertEqual(list(sent.context['exchanges']),[outgoing])
        self.assertNotContains(sent,'Aceitar troca')
        self.assertEqual(self.client.post(f'/trocas/{outgoing.pk}/accept/').status_code,302)
        outgoing.refresh_from_db();self.assertEqual(outgoing.status,'pending')
    def test_admin_can_provision_individual_login_without_changing_imported_identity(self):
        from .models import ResidentAccess
        admin=get_user_model().objects.create_superuser('account-admin',password='test-admin-password')
        imported_username=self.a.username
        self.client.force_login(admin)
        response=self.client.post('/contas/criar/',{'resident':self.a.pk,'username':'residente.ana@example.test','password1':'Temporary-test-only-462!','password2':'Temporary-test-only-462!'})
        self.assertRedirects(response,'/?tab=accounts')
        alias=get_user_model().objects.get(username='residente.ana@example.test')
        self.assertEqual(ResidentAccess.objects.get(account=alias).resident,self.a)
        self.a.refresh_from_db();self.assertEqual(self.a.username,imported_username)
        self.assertFalse(alias.is_staff);self.assertTrue(alias.check_password('Temporary-test-only-462!'))
        self.client.logout();self.assertTrue(self.client.login(username='residente.ana@example.test',password='Temporary-test-only-462!'))
        self.assertEqual(self.client.get('/').context['resident_user'],self.a)
    def test_only_admin_can_create_accounts_and_bad_passwords_are_not_echoed(self):
        self.client.force_login(self.b)
        self.assertEqual(self.client.get('/?tab=accounts').status_code,403)
        self.assertEqual(self.client.post('/contas/criar/',{}).status_code,403)
        admin=get_user_model().objects.create_superuser('account-security-admin',password='test-admin-password')
        self.client.force_login(admin)
        secret='Never-echo-this-password-74!'
        response=self.client.post('/contas/criar/',{'resident':self.a.pk,'username':'bad-account','password1':secret,'password2':'different'})
        self.assertEqual(response.status_code,200)
        self.assertNotContains(response,secret)
        self.assertFalse(get_user_model().objects.filter(username='bad-account').exists())
    def test_duplicate_resident_login_is_rejected_and_own_password_can_change(self):
        from .models import ResidentAccess
        admin=get_user_model().objects.create_superuser('duplicate-admin',password='test-admin-password')
        alias=get_user_model().objects.create_user('already-linked',password='Old-login-password-73!')
        ResidentAccess.objects.create(account=alias,resident=self.a)
        self.client.force_login(admin)
        response=self.client.post('/contas/criar/',{'resident':self.a.pk,'username':'another-alias','password1':'Temporary-test-only-462!','password2':'Temporary-test-only-462!'})
        self.assertEqual(response.status_code,200)
        self.assertFalse(get_user_model().objects.filter(username='another-alias').exists())
        self.client.force_login(alias)
        self.assertEqual(self.client.get('/accounts/password_change/').status_code,200)
        response=self.client.post('/accounts/password_change/',{'old_password':'Old-login-password-73!','new_password1':'New-login-password-94!','new_password2':'New-login-password-94!'})
        self.assertRedirects(response,'/accounts/password_change/done/')
    def test_email_login_ignores_case_and_rejects_ambiguous_email(self):
        user=get_user_model().objects.create_user('person@example.test',email='person@example.test',password='Email-test-password-932!')
        self.assertTrue(self.client.login(username=' PERSON@EXAMPLE.TEST ',password='Email-test-password-932!'))
        self.client.logout()
        get_user_model().objects.create_user('PERSON@example.test',password='Other-test-password-951!')
        self.assertFalse(self.client.login(username='person@example.test',password='Email-test-password-932!'))
    def test_contact_table_parses_three_pairs_and_ignores_notes(self):
        from .contacts import read_contacts
        rows=read_contacts('| Pessoa A | A@EXAMPLE.TEST | Pessoa B | B@EXAMPLE.TEST | Pessoa C | C@EXAMPLE.TEST | Horários em laranja | cobertura UNEX |')
        self.assertEqual(len(rows),3)
        self.assertEqual(rows[0]['email'],'a@example.test')
    def test_bulk_contact_import_preserves_resident_and_existing_admin_password(self):
        from .contacts import contact_preview,apply_contacts
        from .models import ResidentAccess
        admin=get_user_model().objects.create_superuser('existing-admin',password='Existing-password-958!')
        ResidentAccess.objects.create(account=admin,resident=self.a)
        text='ana,NEW.ADMIN@EXAMPLE.TEST\nbia,bia@example.test'
        rows=contact_preview(text)
        self.assertTrue(all(row['ready'] for row in rows))
        previous=self.a.username
        results=apply_contacts(text,{0,1})
        admin.refresh_from_db();self.a.refresh_from_db()
        self.assertEqual(admin.username,'new.admin@example.test')
        self.assertTrue(admin.is_superuser);self.assertTrue(admin.check_password('Existing-password-958!'))
        self.assertEqual(self.a.username,previous)
        fresh=get_user_model().objects.get(username='bia@example.test')
        self.assertFalse(fresh.has_usable_password())
        self.assertEqual(ResidentAccess.objects.get(account=fresh).resident,self.b)
        self.assertEqual(len(apply_contacts(text,{0,1})),2)
        self.assertEqual(ResidentAccess.objects.count(),2)
    def test_contact_import_does_not_guess_invalid_email_names_or_duplicates(self):
        from .contacts import contact_preview
        rows=contact_preview('ana,invalid\nUnknown,x@example.test\nana,a@example.test\nbia,a@example.test\ncarla,pessoaç@example.test')
        self.assertEqual(len(rows),5)
        self.assertTrue(all(not row['ready'] for row in rows))
    def test_contact_import_rejects_existing_other_account_and_ambiguous_names(self):
        from .contacts import contact_preview
        get_user_model().objects.create_user('unrelated',email='taken@example.test')
        self.assertFalse(contact_preview('ana,taken@example.test')[0]['ready'])
        self.a.first_name='José Silva';self.a.save()
        self.b.first_name='Jose Silva';self.b.save()
        self.assertFalse(contact_preview('Jose-Silva,one@example.test')[0]['ready'])
    def test_contact_preview_and_activation_are_admin_only(self):
        self.client.force_login(self.a)
        self.assertEqual(self.client.post('/contas/importar/',{'contacts':'ana,a@example.test'}).status_code,403)
        self.assertEqual(self.client.post(f'/contas/{self.b.pk}/primeiro-acesso/').status_code,403)
    def test_activation_is_single_use_and_does_not_send_email(self):
        from .contacts import apply_contacts
        from django.utils.http import urlsafe_base64_encode
        from django.utils.encoding import force_bytes
        from django.contrib.auth.tokens import default_token_generator
        account=apply_contacts('ana,ana@example.test',{0})[0]
        uid=urlsafe_base64_encode(force_bytes(account.pk));token=default_token_generator.make_token(account)
        url=f'/ativar/{uid}/{token}/'
        response=self.client.get(url)
        self.assertEqual(response.status_code,302)
        response=self.client.post(response.url,{'new_password1':'Resident-chosen-password-462!','new_password2':'Resident-chosen-password-462!'})
        self.assertRedirects(response,'/primeiro-acesso/concluido/')
        account.refresh_from_db();self.assertTrue(account.check_password('Resident-chosen-password-462!'))
        self.assertTrue(self.client.login(username='ANA@EXAMPLE.TEST',password='Resident-chosen-password-462!'))
        self.client.logout();self.assertContains(self.client.get(url),'Link indisponível')
        from django.core import mail
        self.assertEqual(len(mail.outbox),0)
    def test_contact_apply_requires_preview_confirmation_and_no_password_is_set(self):
        from .models import ResidentAccess
        admin=get_user_model().objects.create_superuser('bulk-admin',password='Existing-password-958!')
        self.client.force_login(admin)
        preview=self.client.post('/contas/importar/',{'contacts':'ana,ana@example.test'})
        self.assertEqual(preview.status_code,200)
        self.assertContains(preview,'Aplicar contatos selecionados')
        self.assertFalse(ResidentAccess.objects.exists())
        self.assertNotContains(preview,'data-live-status=')
        applied=self.client.post('/contas/importar/',{'contacts':'ana,ana@example.test','action':'apply','selected':['0']})
        self.assertEqual(applied.status_code,302)
        self.assertFalse(get_user_model().objects.get(username='ana@example.test').has_usable_password())
    def test_existing_legacy_admin_session_survives_email_import(self):
        from .models import ResidentAccess
        admin=get_user_model().objects.create_superuser('legacy-session-admin',password='Existing-password-958!')
        ResidentAccess.objects.create(account=admin,resident=self.a)
        self.client.force_login(admin,backend='django.contrib.auth.backends.ModelBackend')
        response=self.client.post('/contas/importar/',{'contacts':'ana,admin@example.test','action':'apply','selected':['0']})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.client.get('/?tab=accounts').status_code,200)
        admin.refresh_from_db();self.assertEqual(admin.username,'admin@example.test')
